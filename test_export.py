"""测试Excel导出功能"""
import sys
import os
sys.path.insert(0, os.path.dirname(__file__))
os.environ['TCL_LIBRARY'] = r"C:\Users\waiter\AppData\Local\Doubao\User Data\sandbox_runtime\bases\9f6d27f23933fb44a3a1c728c88a5ce4\python\tcl\tcl8.6"
os.environ['TK_LIBRARY'] = r"C:\Users\waiter\AppData\Local\Doubao\User Data\sandbox_runtime\bases\9f6d27f23933fb44a3a1c728c88a5ce4\python\tcl\tk8.6"

from openpyxl import Workbook
from finance_app import _style_header, _auto_width, HEADER_FONT, TITLE_FONT, TOTAL_FONT, TOTAL_FILL, THIN_BORDER, CENTER, RIGHT, LEFT
import finance_app

# 初始化数据库并插入测试数据
finance_app.init_db()
conn = finance_app.get_db()
conn.execute("DELETE FROM invoices")
conn.execute("DELETE FROM salary_payments")
conn.execute("DELETE FROM employees")
conn.execute("INSERT INTO employees(name, monthly_salary) VALUES('张三', 8000)")
conn.execute("INSERT INTO employees(name, monthly_salary) VALUES('李四', 6500)")
eid = conn.execute("SELECT id FROM employees WHERE name='张三'").fetchone()['id']
conn.execute("""INSERT INTO invoices(invoice_number, invoice_date, reimburser_id, amount, purpose, status, created_at)
              VALUES('INV001', '2026-08-01', ?, 234.50, '办公用品', 0, '2026-08-01 10:00:00')""", (eid,))
conn.execute("""INSERT INTO invoices(invoice_number, invoice_date, reimburser_id, amount, purpose, status, created_at)
              VALUES('INV002', '2026-08-15', ?, 1200.00, '差旅费', 1, '2026-08-15 14:00:00')""", (eid,))
conn.execute("""INSERT INTO salary_payments(employee_id, year, month, amount, note, paid_at)
              VALUES(?, 2026, 8, 5000, '半月工资', '2026-08-15 10:00:00')""", (eid,))
conn.commit()
conn.close()

# 测试发票导出逻辑
wb = Workbook()
ws = wb.active
ws.title = "发票报销明细"
ws.merge_cells('A1:H1')
ws['A1'] = "发票报销明细表"
ws['A1'].font = TITLE_FONT
ws['A1'].alignment = CENTER
headers = ['编号', '发票号', '发票日期', '报销人', '金额(元)', '用途', '状态', '创建时间']
for c, h in enumerate(headers, 1):
    ws.cell(row=2, column=c, value=h)
_style_header(ws, 2, len(headers))
conn = finance_app.get_db()
rows = conn.execute("""SELECT i.*, e.name as rname FROM invoices i
                       LEFT JOIN employees e ON i.reimburser_id=e.id ORDER BY i.id""").fetchall()
conn.close()
for idx, r in enumerate(rows, 3):
    ws.cell(row=idx, column=1, value=r['id'])
    ws.cell(row=idx, column=2, value=r['invoice_number'])
    ws.cell(row=idx, column=3, value=r['invoice_date'])
    ws.cell(row=idx, column=4, value=r['rname'])
    ws.cell(row=idx, column=5, value=float(r['amount']))
    ws.cell(row=idx, column=6, value=r['purpose'])
    ws.cell(row=idx, column=7, value='已报销' if r['status']==1 else '未报销')
    ws.cell(row=idx, column=8, value=r['created_at'])
_auto_width(ws, len(headers))

out = os.path.join(os.path.dirname(__file__), "test_export.xlsx")
wb.save(out)
print(f"Excel导出测试成功: {out}")
print(f"文件大小: {os.path.getsize(out)} bytes")

# 验证能重新打开
from openpyxl import load_workbook
wb2 = load_workbook(out)
print(f"Sheet名称: {wb2.sheetnames}")
ws2 = wb2.active
print(f"数据行数: {ws2.max_row}, 列数: {ws2.max_column}")
print(f"表头: {[ws2.cell(row=2, column=c).value for c in range(1, 9)]}")

# 清理
os.remove(out)
conn = finance_app.get_db()
conn.execute("DELETE FROM invoices")
conn.execute("DELETE FROM salary_payments")
conn.execute("DELETE FROM employees")
conn.commit()
conn.close()
print("测试完成，数据已清理")
